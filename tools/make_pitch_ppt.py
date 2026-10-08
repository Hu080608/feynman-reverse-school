# -*- coding: utf-8 -*-
"""生成《费曼反向学校》产品宣传 PPT（16:9，深色科技风）。"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "费曼反向学校产品宣传.pptx")
ICON = os.path.join(ROOT, "tools", "assets", "icon.png")

BG = RGBColor(7, 11, 22)
PANEL = RGBColor(15, 24, 44)
PANEL2 = RGBColor(19, 31, 58)
LINE = RGBColor(42, 58, 104)
TEXT = RGBColor(238, 244, 255)
MUTED = RGBColor(147, 166, 207)
CYAN = RGBColor(34, 211, 238)
PURPLE = RGBColor(109, 140, 255)
GREEN = RGBColor(34, 197, 94)
WARN = RGBColor(245, 158, 11)
WHITE = RGBColor(255, 255, 255)

SW, SH = 13.333, 7.5
FONT = "Microsoft YaHei UI"


def set_run(run, size=18, bold=False, color=TEXT, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    r = run._r
    rPr = r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", font)


def add_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(slide, x, y, w, h, fill=PANEL, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    return sh


def textbox(slide, x, y, w, h, text, size=18, bold=False, color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return tb


def bullets(slide, x, y, w, h, items, size=16, color=TEXT, spacing=10):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = item
        set_run(run, size, False, color)
    return tb


def bg(slide):
    rect(slide, 0, 0, SW, SH, BG, shape=MSO_SHAPE.RECTANGLE)


def accent_bar(slide, x=0.55, y=0.42, w=0.12, h=0.52, color=CYAN):
    rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE)


def header(slide, kicker, title, page=None):
    bg(slide)
    accent_bar(slide)
    textbox(slide, 0.82, 0.38, 11.2, 0.35, kicker, 11, True, CYAN)
    textbox(slide, 0.82, 0.74, 11.5, 0.7, title, 28, True, TEXT)
    if page is not None:
        textbox(slide, 11.95, 6.98, 0.9, 0.3, str(page), 10, False, MUTED, PP_ALIGN.RIGHT)


def footer(slide, text="费曼反向学校 · 教会 AI 才算学会"):
    textbox(slide, 0.82, 6.98, 7.0, 0.3, text, 10, False, MUTED)


def card(slide, x, y, w, h, title, body, accent=PURPLE, title_size=17, body_size=13):
    rect(slide, x, y, w, h, PANEL, LINE)
    rect(slide, x, y, 0.08, h, accent, shape=MSO_SHAPE.RECTANGLE)
    textbox(slide, x + 0.28, y + 0.22, w - 0.48, 0.4, title, title_size, True, TEXT)
    textbox(slide, x + 0.28, y + 0.78, w - 0.48, h - 0.95, body, body_size, False, MUTED)

def build():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    # 1 封面
    s = add_slide(prs); bg(s)
    rect(s, 0, 0, SW, 0.16, CYAN, shape=MSO_SHAPE.RECTANGLE)
    rect(s, 0, 7.34, SW, 0.16, PURPLE, shape=MSO_SHAPE.RECTANGLE)
    if os.path.exists(ICON):
        s.shapes.add_picture(ICON, Inches(0.85), Inches(1.35), height=Inches(1.35))
    textbox(s, 1.05, 3.0, 11.2, 0.5, "费曼反向学校", 44, True, TEXT)
    textbox(s, 1.08, 3.95, 11.2, 0.6, "教会 AI 才算学会", 30, True, CYAN)
    textbox(s, 1.08, 4.72, 11.0, 0.5, "输出倒逼输入，AI 越笨，用户越强。", 19, False, MUTED)
    textbox(s, 1.08, 5.55, 11.0, 0.4, "费曼学习法 × 游戏化通关 × 大模型多模态", 14, False, MUTED)
    footer(s, "产品宣传 · v1.3")

    # 2 问题
    s = add_slide(prs); header(s, "WHY", "为什么我们学了就忘？", 2)
    card(s, 0.82, 1.75, 3.7, 2.25, "只看不输出", "看视频、划重点、抄笔记，都是输入。\n没有输出，大脑不会真正建立连接。", WARN)
    card(s, 4.82, 1.75, 3.7, 2.25, "以为懂了", "看懂别人的讲解，和能自己讲清楚，\n是两种完全不同的能力。", CYAN)
    card(s, 8.82, 1.75, 3.7, 2.25, "缺少反馈", "没有人追问、挑错、要求举例，\n盲点永远藏在“我觉得我会了”里。", PURPLE)
    textbox(s, 0.82, 4.55, 11.7, 0.8, "真正的学会，不是看懂，而是能讲清楚、经得起追问。", 24, True, TEXT, PP_ALIGN.CENTER)
    textbox(s, 0.82, 5.5, 11.7, 0.5, "费曼反向学校把“讲清楚”变成通关条件。", 17, False, MUTED, PP_ALIGN.CENTER)

    # 3 产品定位
    s = add_slide(prs); header(s, "PRODUCT", "费曼反向学校是什么？", 3)
    rect(s, 0.82, 1.65, 11.7, 1.35, PANEL2, LINE)
    textbox(s, 1.12, 1.92, 11.1, 0.9, "AI 不扮演老师，而是扮演一个需要被教会、会主动暴露不懂的学生。", 23, True, CYAN, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    bullets(s, 1.15, 3.35, 11.0, 1.9, [
        "用户学完一个知识点后，必须把这个知识点讲给 AI 听。",
        "AI 会追问、犯错、要求举例；讲不清时明确表示“我还没懂”。",
        "只有 AI 真正学会并通过考试，用户才算通关。",
    ], 18, TEXT, 12)
    textbox(s, 0.82, 5.85, 11.7, 0.5, "把“输出”从可选动作，变成必须完成的游戏关卡。", 17, True, PURPLE, PP_ALIGN.CENTER)

    # 4 机制
    s = add_slide(prs); header(s, "MECHANISM", "三步玩法：讲给 AI 听，把 AI 教会", 4)
    steps = [
        ("01", "用户讲解", "用大白话讲清知识点、原理、例子。", CYAN),
        ("02", "AI 学生", "追问、犯典型错误、要求举例，暴露盲点。", PURPLE),
        ("03", "考试通过", "达到标准，AI 输出“我学会了/考试通过”。", GREEN),
    ]
    x = 0.82
    for i, (num, title, body, color) in enumerate(steps):
        card(s, x, 2.0, 3.7, 3.0, title, body, color, 20, 14)
        textbox(s, x + 0.28, 4.45, 1.0, 0.5, num, 30, True, color)
        if i < 2:
            textbox(s, x + 3.78, 3.15, 0.55, 0.6, "→", 30, True, MUTED, PP_ALIGN.CENTER)
        x += 4.0
    textbox(s, 0.82, 5.6, 11.7, 0.5, "输入只是开始，输出才是真正的学习。", 17, False, MUTED, PP_ALIGN.CENTER)

    # 5 核心优势
    s = add_slide(prs); header(s, "ADVANTAGES", "六个核心优势", 5)
    cards = [
        ("输出倒逼输入", "必须讲出来，输入自然更认真、更主动。", CYAN),
        ("AI 学生机制", "追问、犯错、要例子，持续暴露理解盲点。", PURPLE),
        ("考试式通关", "达到标准才通过，学习结果可感知。", GREEN),
        ("多模态输入", "支持文本、Markdown、LaTeX、图片识图。", WARN),
        ("沉浸式体验", "流式回复、移动端适配、进度可视化。", CYAN),
        ("商业化闭环", "激活码、时长/次数、后台统计一键打通。", PURPLE),
    ]
    x, y = 0.82, 1.65
    for i, (t, b, c) in enumerate(cards):
        card(s, x, y, 3.7, 2.15, t, b, c, 16, 12)
        x += 4.0
        if i == 2:
            x = 0.82; y += 2.35

    # 6 AI 学生行为
    s = add_slide(prs); header(s, "AI BEHAVIOR", "AI 学生如何暴露不懂", 6)
    items = [
        ("主动追问", "每轮至少提出一个问题，逼近知识盲点。", CYAN),
        ("典型错误", "故意犯该知识点的常见误区，等用户纠正。", WARN),
        ("要求举例", "要求用户举例、换种说法、做小题。", PURPLE),
        ("明确没懂", "讲不清就说“我还没懂，因为……”，绝不敷衍。", GREEN),
    ]
    x = 0.82
    for t, b, c in items:
        card(s, x, 2.05, 2.85, 2.7, t, b, c, 17, 13)
        x += 3.05
    textbox(s, 0.82, 5.25, 11.7, 0.6, "AI 越会暴露不懂，用户越必须真正学懂。", 22, True, TEXT, PP_ALIGN.CENTER)

    # 7 功能亮点
    s = add_slide(prs); header(s, "FEATURES", "学习体验亮点", 7)
    features = [
        ("Markdown + LaTeX", "AI 回复支持代码块、表格、公式，适合数学、物理、计算机等学科。", CYAN),
        ("DeepSeek 识图", "图片直接识别文字和公式，自动整理为 Markdown / LaTeX。", PURPLE),
        ("逐字流式回复", "像真实对话一样逐字出现，学习节奏更自然。", GREEN),
        ("账号与余额可恢复", "换设备可用激活码恢复已有账号，余额不丢。", WARN),
    ]
    x, y = 0.82, 1.65
    for i, (t, b, c) in enumerate(features):
        card(s, x, y, 5.65, 1.85, t, b, c, 17, 13)
        if i % 2 == 1:
            x = 0.82; y += 2.05
        else:
            x = 6.85
    textbox(s, 0.82, 5.85, 11.7, 0.5, "从输入、理解到输出，所有环节都被重新设计。", 16, False, MUTED, PP_ALIGN.CENTER)

    # 8 商业化
    s = add_slide(prs); header(s, "BUSINESS", "商业化与授权机制", 8)
    card(s, 0.82, 1.65, 3.7, 2.4, "一次性激活码", "服务端 HMAC 核销，一码全平台仅兑换一次。", CYAN)
    card(s, 4.82, 1.65, 3.7, 2.4, "时长 / 次数套餐", "支持按小时、按天、按次数；可加时、可叠加。", PURPLE)
    card(s, 8.82, 1.65, 3.7, 2.4, "暂停与续接", "进入聊天页计时、离开自动暂停；历史对话可恢复。", GREEN)
    card(s, 0.82, 4.35, 11.7, 1.7, "公众号分发闭环", "用户在公众号下单 → 你生成激活码 → 发送给用户 → 网页输入激活码 → 后端核销并记录。", WARN, 18, 14)

    # 9 技术架构
    s = add_slide(prs); header(s, "TECH", "技术架构：安全、稳定、易部署", 9)
    boxes = [
        ("用户浏览器", "GitHub Pages 静态前端\n聊天记录存本地", CYAN),
        ("Cloudflare Worker", "鉴权、套餐、限流\n隐藏 DeepSeek API Key", PURPLE),
        ("DeepSeek API", "文本对话 + 视觉识图\n流式输出", GREEN),
        ("KV / 运营后台", "激活码核销、余额\nToken 用量、金额统计", WARN),
    ]
    x = 0.82
    for i, (t, b, c) in enumerate(boxes):
        card(s, x, 2.2, 2.75, 2.6, t, b, c, 17, 12)
        if i < 3:
            textbox(s, x + 2.78, 3.35, 0.5, 0.5, "→", 26, True, MUTED, PP_ALIGN.CENTER)
        x += 3.1
    textbox(s, 0.82, 5.35, 11.7, 0.8, "前端不暴露大模型 Key；授权、计时、用量统计全部在服务端完成。", 17, True, TEXT, PP_ALIGN.CENTER)

    # 10 目标用户
    s = add_slide(prs); header(s, "USERS", "谁最适合使用？", 10)
    users = [
        ("学生 / 备考党", "数学、物理、计算机、考研、雅思托福。", CYAN),
        ("知识工作者", "快速吃透新领域、写文章前先讲一遍。", PURPLE),
        ("教师 / 培训师", "用“反向教学”检验自己是否真懂。", GREEN),
        ("自学者", "读书、看课后，把知识讲给 AI 学生听。", WARN),
    ]
    x = 0.82
    for t, b, c in users:
        card(s, x, 2.15, 2.85, 2.6, t, b, c, 17, 13)
        x += 3.05
    textbox(s, 0.82, 5.3, 11.7, 0.7, "只要你需要“真正学会”，而不是“看起来学过”。", 20, True, TEXT, PP_ALIGN.CENTER)

    # 11 商业模式
    s = add_slide(prs); header(s, "MODEL", "商业模式与增长路径", 11)
    bullets(s, 1.05, 1.8, 5.3, 3.6, [
        "① 体验装：1 小时 / 少量次数，低价引流。",
        "② 标准装：按天 / 按月，满足持续学习。",
        "③ 次数包：随用随扣，适合低频用户。",
        "④ 加时包：有效期内叠加，提高复购。",
        "⑤ 后续：题库、学习报告、班级版、企业培训。",
    ], 17, TEXT, 14)
    card(s, 6.8, 1.8, 5.7, 3.6, "公众号卖点", "“教会 AI 才算学会”\n\n不是卖课程，而是卖一种\n可验证、可通关的学习结果。\n\n用户越用越强，越强越想分享。", CYAN, 20, 15)

    # 12 路线图
    s = add_slide(prs); header(s, "ROADMAP", "路线图", 12)
    roadmap = [
        ("v1.3 已完成", "AI 学生机制、流式对话、激活码、计时/次数、后台统计、DeepSeek 识图。", GREEN),
        ("v1.4 计划", "学习报告、薄弱点分析、分享战绩、自动发码。", CYAN),
        ("v2.0 计划", "题库训练、班级/小组、教师后台、多模型接入。", PURPLE),
    ]
    y = 1.7
    for t, b, c in roadmap:
        rect(s, 0.82, y, 11.7, 1.25, PANEL, LINE)
        rect(s, 0.82, y, 0.1, 1.25, c, shape=MSO_SHAPE.RECTANGLE)
        textbox(s, 1.12, y + 0.18, 3.0, 0.4, t, 18, True, c)
        textbox(s, 4.0, y + 0.22, 8.2, 0.8, b, 14, False, MUTED)
        y += 1.55

    # 13 结尾
    s = add_slide(prs); bg(s)
    rect(s, 0, 0, SW, 0.16, PURPLE, shape=MSO_SHAPE.RECTANGLE)
    rect(s, 0, 7.34, SW, 0.16, CYAN, shape=MSO_SHAPE.RECTANGLE)
    if os.path.exists(ICON):
        s.shapes.add_picture(ICON, Inches(5.95), Inches(1.05), height=Inches(1.1))
    textbox(s, 0.8, 2.65, 11.73, 1.0, "教会 AI，才算学会。", 40, True, TEXT, PP_ALIGN.CENTER)
    textbox(s, 0.8, 3.85, 11.73, 0.6, "输出倒逼输入，AI 越笨，用户越强。", 21, False, CYAN, PP_ALIGN.CENTER)
    textbox(s, 0.8, 5.0, 11.73, 0.5, "体验地址：https://hu080608.github.io/feynman-reverse-school/", 15, False, MUTED, PP_ALIGN.CENTER)
    textbox(s, 0.8, 5.6, 11.73, 0.4, "制作者：胡胜杰 · v1.3", 13, False, MUTED, PP_ALIGN.CENTER)

    prs.save(OUT)
    print("已生成：" + OUT)


if __name__ == "__main__":
    build()
